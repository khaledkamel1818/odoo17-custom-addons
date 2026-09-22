# -*- coding: utf-8 -*-

import base64
import csv
import io
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError


class DmFinanceBankReconciliation(models.Model):
    _name = 'dm.finance.bank.reconciliation'
    _description = 'Bank Reconciliation Workspace'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'date desc, id desc'
    _check_company_auto = True

    name = fields.Char(string='Reference', required=True, default=lambda self: _('New'), tracking=True)
    date = fields.Date(string='Statement Date', default=fields.Date.context_today, required=True, tracking=True)
    company_id = fields.Many2one('res.company', string='Company', default=lambda self: self.env.company, required=True)
    journal_id = fields.Many2one(
        'account.journal',
        string='Bank/Cash Journal',
        required=True,
        domain="[('company_id', '=', company_id), ('type', 'in', ('bank', 'cash'))]",
        tracking=True,
    )
    branch_id = fields.Many2one('dm.finance.branch', string='Branch', domain="[('company_id', '=', company_id)]")
    state = fields.Selection(
        selection=[
            ('draft', 'Draft'),
            ('imported', 'Imported'),
            ('matching', 'Matching'),
            ('confirmed', 'Confirmed'),
            ('cancelled', 'Cancelled'),
        ],
        string='Status',
        default='draft',
        tracking=True,
    )
    import_file = fields.Binary(string='Bank Statement File')
    import_filename = fields.Char(string='Filename')
    line_ids = fields.One2many('dm.finance.bank.reconciliation.line', 'reconciliation_id', string='Statement Lines')
    line_count = fields.Integer(string='Lines', compute='_compute_counts')
    matched_count = fields.Integer(string='Matched Lines', compute='_compute_counts')
    unmatched_count = fields.Integer(string='Unmatched Lines', compute='_compute_counts')
    note = fields.Text(string='Notes')

    @api.depends('line_ids.state')
    def _compute_counts(self):
        for rec in self:
            rec.line_count = len(rec.line_ids)
            rec.matched_count = len(rec.line_ids.filtered(lambda line: line.state == 'matched'))
            rec.unmatched_count = len(rec.line_ids.filtered(lambda line: line.state in ('unmatched', 'suggested')))

    @api.model_create_multi
    def create(self, vals_list):
        seq = self.env['ir.sequence']
        for vals in vals_list:
            if vals.get('name', _('New')) == _('New'):
                vals['name'] = seq.next_by_code('dm.finance.bank.reconciliation') or _('New')
        return super().create(vals_list)

    def _check_reconciliation_operator(self):
        """Keep write operations protected even when called through RPC."""
        if not (
            self.env.user.has_group('dm_finance_enhancement.group_bank_reconciliation_officer')
            or self.env.user.has_group('dm_finance_enhancement.group_finance_manager')
        ):
            raise AccessError(_('Only bank reconciliation officers or finance managers can modify a reconciliation.'))

    def _parse_date(self, value):
        if not value:
            return fields.Date.context_today(self)
        value = str(value).strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y'):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        return fields.Date.to_date(value)

    def _parse_amount(self, value):
        if value in (None, ''):
            return 0.0
        cleaned = str(value).replace(',', '').strip()
        return float(cleaned or 0.0)

    def _iter_import_rows(self):
        self.ensure_one()
        if not self.import_file:
            raise UserError(_('Please upload a CSV or XLSX bank statement file.'))
        filename = (self.import_filename or '').lower()
        content = base64.b64decode(self.import_file)
        if filename.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
            except Exception as exc:
                raise UserError(_('openpyxl is required to import XLSX files: %s') % exc)
            workbook = load_workbook(io.BytesIO(content), data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(cell or '').strip().lower() for cell in rows[0]]
            for row in rows[1:]:
                yield dict(zip(headers, row))
        else:
            stream = io.StringIO(content.decode('utf-8-sig'))
            reader = csv.DictReader(stream)
            for row in reader:
                yield {str(key or '').strip().lower(): value for key, value in row.items()}

    def action_import_statement(self):
        self._check_reconciliation_operator()
        for rec in self:
            commands = [(5, 0, 0)]
            for row in rec._iter_import_rows():
                amount = rec._parse_amount(row.get('amount') or row.get('balance') or row.get('قيمة') or 0.0)
                commands.append((0, 0, {
                    'date': rec._parse_date(row.get('date') or row.get('التاريخ')),
                    'amount': amount,
                    'reference': row.get('reference') or row.get('ref') or row.get('المرجع') or '',
                    'partner_name': row.get('partner') or row.get('customer') or row.get('supplier') or row.get('الشريك') or '',
                    'memo': row.get('memo') or row.get('description') or row.get('البيان') or '',
                    'state': 'unmatched',
                }))
            rec.line_ids = commands
            rec.state = 'imported'
            rec.message_post(body=_('Bank statement imported.'))
            self.env['dm.finance.audit.log'].log_event('bank_reconciliation', rec.name, rec, _('Bank statement imported.'))
        return True

    def action_auto_match(self):
        self._check_reconciliation_operator()
        for rec in self:
            for line in rec.line_ids.filtered(lambda item: item.state in ('unmatched', 'suggested')):
                line._suggest_match()
            rec.state = 'matching'
            rec.message_post(body=_('Automatic matching suggestions generated.'))
        return True

    def action_confirm_match(self):
        self._check_reconciliation_operator()
        for rec in self:
            suggested = rec.line_ids.filtered(lambda line: line.state == 'suggested' and line.suggested_move_line_id)
            suggested.write({'state': 'matched'})
            rec.state = 'confirmed'
            rec.message_post(body=_('Suggested matches confirmed.'))
            self.env['dm.finance.audit.log'].log_event('bank_reconciliation', rec.name, rec, _('Bank reconciliation confirmed.'))
        return True

    def action_reset(self):
        self._check_reconciliation_operator()
        for rec in self:
            rec.line_ids.write({'state': 'unmatched', 'suggested_move_line_id': False, 'match_score': 0.0})
            rec.state = 'imported'
        return True


class DmFinanceBankReconciliationLine(models.Model):
    _name = 'dm.finance.bank.reconciliation.line'
    _description = 'Bank Reconciliation Line'
    _order = 'date, id'

    reconciliation_id = fields.Many2one('dm.finance.bank.reconciliation', required=True, ondelete='cascade')
    company_id = fields.Many2one(related='reconciliation_id.company_id', store=True)
    journal_id = fields.Many2one(related='reconciliation_id.journal_id', store=True)
    date = fields.Date(string='Date', required=True)
    amount = fields.Monetary(string='Amount', currency_field='currency_id', required=True)
    currency_id = fields.Many2one(related='company_id.currency_id')
    reference = fields.Char(string='Reference')
    partner_name = fields.Char(string='Partner')
    memo = fields.Char(string='Memo')
    state = fields.Selection(
        selection=[
            ('unmatched', 'Unmatched'),
            ('suggested', 'Suggested'),
            ('matched', 'Matched'),
            ('ignored', 'Ignored'),
        ],
        string='Status',
        default='unmatched',
        index=True,
    )
    suggested_move_line_id = fields.Many2one('account.move.line', string='Suggested Journal Item')
    match_score = fields.Float(string='Match Score')

    def _suggest_match(self):
        self.ensure_one()
        account = self.journal_id.default_account_id
        if not account:
            return False
        amount_domain = [
            ('company_id', '=', self.company_id.id),
            ('parent_state', '=', 'posted'),
            ('account_id', '=', account.id),
            ('balance', '=', self.amount),
            ('reconciled', '=', False),
        ]
        candidates = self.env['account.move.line'].search(amount_domain, limit=20, order='date desc, id desc')
        best_line = self.env['account.move.line']
        best_score = 0
        for move_line in candidates:
            score = 50
            if move_line.date == self.date:
                score += 20
            if self.reference and self.reference in (move_line.ref or move_line.move_name or ''):
                score += 20
            if self.partner_name and move_line.partner_id and self.partner_name.lower() in move_line.partner_id.name.lower():
                score += 10
            if score > best_score:
                best_line = move_line
                best_score = score
        if best_line:
            self.write({
                'suggested_move_line_id': best_line.id,
                'match_score': best_score,
                'state': 'suggested',
            })
        return bool(best_line)

    def action_mark_matched(self):
        self.reconciliation_id._check_reconciliation_operator()
        self.write({'state': 'matched'})

    def action_ignore(self):
        self.reconciliation_id._check_reconciliation_operator()
        self.write({'state': 'ignored'})

    def action_reset(self):
        self.reconciliation_id._check_reconciliation_operator()
        self.write({'state': 'unmatched', 'suggested_move_line_id': False, 'match_score': 0.0})
